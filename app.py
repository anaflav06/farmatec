import streamlit as st
import pandas as pd
import json, os, re, hashlib, io, base64, zlib, requests, unicodedata
from datetime import datetime, date
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

st.set_page_config(page_title="FARMATEC | Controle de Saldo", page_icon="💊", layout="wide")
CLIENTE="FARMATEC"
CNPJ="06177615000174"
CUTOFF=date(2026,9,22)
DB_FILE="farmatec_database.json"
REMOTE_DB_DEFAULT="data/farmatec_database.json"

# Segurança e persistência para Streamlit Cloud/GitHub.
# Configure em Settings > Secrets do Streamlit:
# APP_PASSWORD = "SUA_SENHA"
# GITHUB_TOKEN = "github_pat_..."          # token com permissão Contents: Read and write
# GITHUB_REPO = "usuario/repositorio"      # ex.: gdslogistica/farmatec
# GITHUB_BRANCH = "main"
# GITHUB_DB_PATH = "data/farmatec_database.json"

def secret(name, default=""):
    try:
        return str(st.secrets.get(name, default)).strip()
    except Exception:
        return str(os.getenv(name, default)).strip()

APP_PASSWORD=secret("APP_PASSWORD")
GH_TOKEN=secret("GITHUB_TOKEN")
GH_REPO=secret("GITHUB_REPO")
GH_BRANCH=secret("GITHUB_BRANCH","main") or "main"
GH_DB_PATH=secret("GITHUB_DB_PATH",REMOTE_DB_DEFAULT) or REMOTE_DB_DEFAULT

def require_login():
    if st.session_state.get("farmatec_logged_in"):
        return
    st.markdown("## 🔐 FARMATEC — Acesso restrito")
    if not APP_PASSWORD:
        st.error("Senha do app ainda não configurada. Cadastre APP_PASSWORD nos Secrets do Streamlit.")
        st.stop()
    pwd=st.text_input("Senha", type="password", key="farmatec_password")
    if st.button("Entrar", type="primary"):
        if pwd == APP_PASSWORD:
            st.session_state["farmatec_logged_in"]=True
            st.rerun()
        else:
            st.error("Senha incorreta.")
    st.stop()

require_login()

BASE_BLOB="eNrVfduO40iS5a8Q9TzD8ftFby7JQ8FMilSRVGxFLAaDnOoCphbd1Y3eHszDYP59zZyKCJGSewKk4ImtKiCz8iKJJnezY2bHjv3v//7p23/9+0+b4ieptTLcECvET/9U/PSnb//4DX+ZEab+mdh/pnIgdsPNhotSaU0IwT/1f3/740+//f3ffv3jb/8H/7Ci8CJUwm9SLa5+/49vfwkvNvzHt+LbX/5WDP/xe/H07e9/+fbr79+KP/1WHL/98fvf/vPP33799lf8W3//7dff//b7b3/849/+9Ndfxw+hmFT4poSOb/z5R95fu3fnvS+2/lj4fnBd0fvupdq1fXH0+/BjPewd/tV//PUf3/4Mf0Pakmv8jH/+Fp4/fN6//ufffw0vt3W9L6qm2lWuLhj7F2L/BU3x0//8UzE3GSU0ZTK9EaQUjOc1mbRMSkKptvfN9cX3fbVzxbHqXLN3Bdiub89vrqh9NfhrO1HKS0I/DaUXGYpQprnU9w1lBqo2lGwEKxV+yVnPljSKc22FuW+ouqor1xT1efda7OEY9a4Z2n5iIMJKYT8NpBYZyFpuORUkZiCxYXJDdClyXz4pmdRwlpTR9y3ktp17dm1RV82ubWqwVTsMxeCrXzwcLjxZbV294M+nx4qUXKy1Gh4ryThPHCu7YbYkIrfLMuCxuFEy4q5c0/ji2TXOF0fXD/352A7ttX0UL436NI9YeKiUBr+p4oeK2w2RpZGZzUMlpxxcJ5f3zXPy+64tnn3TVT+fffTi6VLatZ7JWi3xk8RtJCg4p5Lq3C7cWC4h8OELs4iddjU+YLhkT1VftYPv3KlyxQEuZeVruIV17e/EPl1SsfZwhbunjYnfPWY3wvwAuMAoVTIEEspUBC4M533VFqfiMDXc3FIQQUsuP01llpvKkISpONkwVQYvldNUShmjVIBWVNDIXWzr+hU8VQEwoa/qF1dsz76ZOitb6vUB0HBiZeoe8g0zJaeZbWS1hauoAMKJ+wba+wqe8Ogb/wZ37cl3XQh8/e75fHS7Z99NDhSgQU0fcKC4EJTFjSXFhprSstxwigrIIKwKLxwBDDvfDODdR2Be7NwX9xYMNr97RpSGr0foxCihZByhU7phvKTZgSeDZEYARtARiACQs8Jr17X7rjqc/RgC8Vf61nV+EgoVK8lqCKoMUwRCTtxUDC4heCmVORJSLeGlpY0eqK7qB7CVL9wJLLOrwFO9Y87i5Ls59tSspGvPlaFWcM7v43WFxpJkw2VJs2d+cNiV5CH6UZqGDe7VFe+Rr2oGf+jcUL3c3kQFX7pcixcMo4QzKRIW4xvKSq4y30QuKBVGGAiFEbDw7KsaQBXkxjt/AaJjunxtJKsnBQW11EjUSK4iRoIkGRE7hXfK7a4A6XGI0Ybw+0Y6ug4sU2x977tiD0/69gEYZmUXs/7yMaqEIDRxlPQG0nKZ21NZSQyXlLBYmaqu9sFR1dXRFcf2PE+K6TWQ0kuNY4RiJH6ECMVKgiK5KwkC0IGgPJTxIletasCTD+ehLZ7cGX7atIWHf8Foh6aF/79XyaN84s+XWg3isVI8bjXKIROENDDzxdPUQuqp5aU2dMeZuw6CHcACQJ6Na/b+bRLrbKnU+vsm4ME0jRuHgXF4CTlOZtdNpBRacBNLXp7cFtK8+gze6IhZMqTF0zoLJniKfRrILjWQUoyl3LbYcF2S7IUWDl+JUdrEsBO47QNgpsGjdVx99NX0cmmAMHb9+ZGSqpR5wF8DsrQ2s3ngXhlBDDUxd+3OEMp8hSi82bfN3ZCGKR1T64+QNFqx6B2DLIViK4HRzDFNSWYRgMtIK+GrCzW6g6v36K7BHe3bfnrHLCu1We+hJWThPFohZxumN1Lld0JWacIFp0yzSLOlxaztq+/hHEEAG9xdT8QYhDG59hhZxhnnkXKmHqjcEL5hpFQye0UcYMkIW6n4Tjlz544nCPT4AwS2O/VLVSq9ujiuAc0qIWPniYaOC8BVkxsnEY6ZUnjdSFcq1E4w5r+4fneuXXebsomSrI77kHHDh9EqYSIaanIy85UzFsI++CXB0u3NfdVCIlJNi5WkNGK1ZeDoEBnx1sEZCQBFkKfldkYQQhikICoFiEIx1xcntz87iPkvvi+6auurrr3bqpNmEt4W3jZjEfZHG8B8Q2Tw3rmPkpSMUz1WLE3EL3Xt1nfgtl+qUFLy/eAHPF5YDEDawaV2sr/TP9CqvKph8sXVXmlU1FNJtB2kjmji3PDbQtwzqWpvf3K79mKn3XFqHGJK+YgesIQDb+LWAVxAIWLQzH7cgheXAr4THkGXT7V7QWjpD+jKd23X+WkBnJdiNWySWFHiIloAN1h247S0IvfFA9zE5XjxeKSo5MfTA0DgtQafdXYdJCo3d4zJUqztqCiDLXue6BNwFlrlufN/yq2E9DJKfGrc4JoK+7wF5HJwnAAyPQdfHmFeKFNSvb6YKwQglERGh2Bcljp3AwrCHwU/GNCTVLGKybGq6/du5lPnqwFAZsRcyBOTq8u6VFlIMxO1bwaXELKk7D7KEkkZUTyWu+zD4XLvVe9t2x2mrSe4u6Xgq1kYgMSNEDyOpwjZQLjI3hyQSO+TlLPY7Xt1gAje4dOpBSTQz26btuvBkxIavHgch8tAc8runKQggIRZuG2x7gl4pZ2v22Joj7s393UKme62mwDtQNa11qOD0aiVMpHfCY1FS50dNVEhCGYQMZfenSHKfTY137mZMwbPauAUDERE4lQJtRntlyvWMc2FEEarSK2pc08OvdFHd7evjm1IX6bUMDCPtutzPEWZNYnsV8B/Mn+BQGjLjNVIWaAqcoZeqmFw412bems5CWeLvTWll+5yxDIIAEqSO/pzJggzYB6dZPNW9dfqPOnd2tLa9UFeglmsigd5wrArabM3ARjSvCiAtqiLPnfwdBC/nsFAp5Mvtm737FwNwX9WCTDrrcSEZjYKhcxGsA23pWQ6e+9WQDJpaAwKndyAVnJgnG7b9q44tMdZsFclfUDRDXKR2DEKlwup8qLkNHuzBJI1GjiK0YStenP7Fvzy0e38PqQgXbsf2UsHHMxoC3+845WQ0bC69E3QamysQ8wNx1hI4dSGwFeU3StZbRBlm9iQwZcWq9zIcoZcpC+aFgw2S0OMmNAG9XILKWYTFkJ2AOSHmY+WhgxEaqUpi4Qzj7n//poM98FGbc7N9BYaO6kqycWmspfoGjEVJLg0Py0cUkRKFIPslseaKJem5RVtHvI29+KbYc434aqUVw6L8uW2gt+I24ppHBwTuacwACHB18OJ5RE8cKFMuOILuPJumDAnjC756iREa4psWBopK1EVMKSBDDp3V5czCxAyNljwv3y9913x1Da9xz5KW4CH2lVHOEMz+oQs7VU1aaQkLzGTNPHqG5oJXDe4QJHddTMDKUm0nfLlXL+3U+AcdQ4Dnn/zDeZs/uC2VePe3iYYyoLPYOurldQwYhNUXRx9YiXPnZpAqBPwsqGoa3Vy9mKsdxdbX/u3z1bKDQkcMrnVrErFOKU8Un4TA9Ub9FGAQXJjTkoNchljvFNIbpsPwmmkPmn5BHMuNZDU8vI57hsI7h+AcqryQyfOGIuV37r2AF4Kb54vjiMLDsJfN5sTgxzdrvbkiilInqmMm0hS7KPI3BwmK7gi4BjBI8TQZe/ro7vy4Jem72WeYEL3UiWT60+TohByecJUEsuVTGQOehx5lERCMhypL/mRn1vsXF39MrQxThykwavnDsFITFlCYkYyWPFmppQ6e1MX7puylETc9zMY6ZMogDg8cJrmk5nXXtsuBQUQQ2QCFHCBEBwHx/JiJ4osZiFJJJ2Dw4NlgsDxRgzumi9tn5inkEvNA7lcpCkQzAMOieKkZ+6mgLIQb1VgBbBIknJ2wVWPdK8b7pJ+AENHY6IkmEqYh+OATnZ/jYc63Gms3kZc9vDsKuS/1Xs3ThS6t6vBwnkzQJFrjolabC5IeBOJijRYGhC55+S4ApBkxWXQ+G57qaoBQHbnp0A/bWscjrs/e2ImwW3xuTKUWJI0FCCO7PVLSHalpRyyTRurX9YteKSxsQTX71B0bVe9TR03KdXayIYWQtmBmIUAbStESiHlzUtB5cQaTbmKnKXG13BofBOoFGAjuIKtr5/b4gu2vedjA5Ktj3AMUiWasBQl48yzzs2i4AISE85j9bj9J+fk4OpnBzdu21VTG0FmItV6vwR5rjUyYSK+YaLUuWESpLnYgLM6kry9uMb3vbsaST24DmEB1uHAZT1P6yjETsoCS28eo0wrnjAWcgXzy6HA+TWSMxqbdXb7LpRR3rk4vTueXDXXGDCrNSvAQFzHaLujgcwmXO3ctF24aVwQHRtEcbX/BWe7ugsboEelpm3n2jmf0q4ulsD3RDSL2EgE2E2Rq8QWlyyfXVEMz1Xx5LqjwyuxRz/SVKdz7XaoJHTBg7duyVhG5cg3jZV0lSngDxXwp4pd1QzPr5c+XTFgvamfzYGz1e06MBe3JCLFEMwlwWKQxy0GTovNxbEobxQhsXjX7nHwa/Bgl6qBtzm13eDr+fyXWD3bhKNxnMaPFA7rbhj7AWxBqZAuaMNIaqQU/nKsi5Pre0jkfLG7oxHDBCmNXt1F0cpobqJEXeQxG4RNOnerQGorqOBCstgAWFe74ovbffWQ0Pli+4Ge7k7LQRLB9HpYANHExrM7jeQuDslv7mk5LiiXkP0yImKjOxjq8DJXEOjGQtx9fSvNJvhpYcSDZFwIpeKnSjAcvTTZm+KSwd0T1JLYJLhDRsqXc1OFOtwFGkwgAQUks/okKSuFICJuIBl64iS3kIdgqDFCRIw0sPdbZFoMvoEsrh3J35AA31w4bSdkSr3YStYwkrASismV1mb34MEzBQ8eq+peJuRfi6+YsjSXSwdPfiMoZx8wcAGZCNeWJcqW4b/S5tZf0FQbYa2QsV6m2/UtWsf/UjSICYoL6uymqS8jJeXrG+Mk8K3iVmIEIQHLrmDFhCXgMk2sjBIMcsUCi6icXLN1liYrgMM5TZlIIrLM3kah3FCmUHdBRsuXr1eCFH3RueMstqlSrdb2Qlk9oqWNy+rhrDMvkQKS10ACx3cZVTzFZmqLU2jH+TrFrzClkY8wFHxjKUPJkPZm118g4Cvhk6WUnfcusOCDJGjf3r1uONKkHiFRLC9KEBEtWRM4FTJ7gxcyFCZ1bIxwcJc5gU+618G9VBDy5jrOnD3ASIgURYzqZS6TKNlru0HnhHER5r7irYKgCH7JTibCn50/eAh69wdSpCmtfIBeqgS0ouJzqhx7B6XNrlWswZVThROQsVmCl3DA3lyze/bDMJ9EUfQBDgpso2nCNhrvHqHk/1/VCi0eoFqBljI2PgomkalDAX7kRpdGGaGoIJTKeKNupOd8YKZRoLHthvaevDMlE1315SfLEiMS9sJ+VClz83jhZCnwpCFxoSaWA+/btiv88QRofCQZ9n3VD6N2KpILjvD8dXVH/8PQSX9qKQJFiVIrEgwM1EUDl5WbgSGM0UKOusUsfS0vk4bdhWDoP+UbbngHdMJakUszZSMBlcZtRkYV1exTQIIAkBgz5ZiIY3V8+VTg2wdOnccpxP5DaOaemjgrJX3AMCI3hidGyiTfwDup7HQohctHJIkNk5+6qt+hOkG/AwfXDlN+PZm4/JFescg2WqaUisA2qB5lM19DwwTm59amARfgq7c3xA13lMPFZHhcLm0vMMltomN1YUJlRw+A24WGc01Yukm8c3DLZi1hZKvw1ToNKIYgOEvIWtqNFPk9OFwrAe8ZPLiOsqDao3u7Jj5fRn7gp36a3WhSygdU7hgHB8mT3WFs9uTuwAiBDTQhTax0B1kMFsvDgNTuOUBR1x0AVbXbrmpm50o9IsppBiGYJ2gZjG8onKvcpmJWUaPi2h8utKhCaDuPXLp+JsHLSiYeQOyRjLKUeTRgp9Ky7NfOSiSNpuYzxiH74lC5Xbvdtt9fpaUpPMgDqBkA6QSLmwxSQA4AILemM0cGkhj1+GhsB9KxfceXYQzoc5bz5DpUxzy64/a8c9t7G0bAenR9OZ0L1MSOWw/hE7xT7gEO7K5zwWWMSFbVgW4fNEPv94vxmjyiAir1pcUfSQPphqv8xT2AUAAtIXOWPJU2f+xnG9zLfPkDFQb87HqNXk6EjrisUTZcbAgpqcpdIw6BJqQtsVYoEn2qU+d+qfq76mD6AQrGHGXUUjhK/ZCsDhVTiJWjFkhUVN13p+JqV1Tjz11bt4fxYI3iRa4tTp0f2nuLEtfvFkH7GWVT9iOhbJW7+MkNZk1jByJG5Wx8CyjdN7vb8URGp2hqqbQxOAEpE/rq46xC9uQXHSfkEJRZGisN1zU4bzhR2FI/ovxV1d0IrJvVA6/wLVmrRUJbRWg8Pyz7Bi0tiWDcpLa0vVSDO1ZNgQ2Gw+uthxJTTZ6x47vgECGXXMi4gDhkL/wHCIhTYoyxWMIgaU3aS/Pq6Ophvq/nITwfTbVMjEvhvB35ARmLgU8h4wjJ98Oz78AqryF1OfrQUUek+QaXbt/O2Jrr10ABCueMibhgqEDxi1Ll7qpDDmy4Seim9SP2/lg9uu3cYb7HiK8uqUgqteDRASC0T+jAKJ5dUJVbZQmnMWf9pQ1zwGMN5SJZOBnYEKWWD+BDA9RWJF6MExKZYtzm50NTInDNbugRR2y0b48F+us6KkiAIGX9hDRVMpauXRsp+1QLA09pFTOUxTLdoJfyOdVydIfm3CNxLAyW+WbWHOarW554oqyhKlne5XQFD2rxiQLQpiAxiAX+IPCM42SHttldqL735xKlXV8E55THenXiMkuGDZTcjWHcYiu5trHaybuE6kVjHfLduip2uFW6nudxVK6fuuc4355QcUCePfZOc5d4gz4Y08aSBO3wtUAxB7AOXL+uffso+M4c1HWQs0vvnNHakLidOEF9Z764GrD4zlnCNNL6Waxi4uvqzW398Fw8j3zfj4nXT8HnKaXViAeUmCTH/e2SJTRULJbEhczvpqwlVgmbara8VC++2/pjsWvb2m1vm+NssidrMaWVKisTU+YBPpUid6dO4EQZDWO2lEUmN7benYfX4qntjv31Dog7GzSnenRqsoJt4ZYDiIEsoqkyaqxRjo08mRt1agGwnGlhVWxM4RU9Vu0D1TUUMm9lZ5CceM28t4ttxI0UCR06rGKWNLeNDCAYlDejsRXSmP3iGN7bx5QwhMTzF7jSM7UnCtfwutG5WLEPd+eolGIfwxwvu7SKRglPMbanotXMJ3hAQAojF8o3aK4LIcoXz21/Qvh+byWSnsg/iMWWg/OeuIci7InMXsiU2hIB/8ALs1ih5UNw5Yxp4Htbb+9vN9rc1KjsVOdvcW4otBAJ9EUgdSYrulSLQ6Rg8LKUURMT9Tk5eIkrKDHWqgKon2U8xDygW6WMNDbercL9tqJk2VeTUdStx1XJ32lWPUEYPH50q/BXkEAGWZCvZ1JI9hGXEocbSIo+jPslWSl0duKGZTRUrOLzxcc23Ml+GNdspPdJaT1pFi+VAQakKu8vbwEfhmvdAE6Y/JsRLGSK+KqMktiI2s/nccj4Mj/7Mwolt1MNGzldcCMX20hSTeI24gY1bbPLsxlIMhQVKBcRcfLvLgqyRI/dmWouzbaeKRWsI5mIW0ewjbT5haSlsIZeusU0AiCeiv27GPkFNtxyMiykIOshOxjJcJMwUpDU0rlrDRTZW+CX6HfUWJ79uZ/r/1OxHkGhnJGxcbNINbLscnsfVIpmRMc22YW1Ua44de0LgE6ctQ79zrq+WdpmH+B9mGT36SpoI4ZAScj8eTKn2CZnLNpjwPLdKxhmlKr7qHvORhMEfcACMgUhNbHlHgtUkOuZ7LOfioKL5jKqMTZ45GpelYTfuT1B66CbxjL1gE1I8H1JRrVMFND1hvP8/GjNmNIsusPmiw+Dnx9Dsh6y4u6M3Qc4X9VxHtaur51ZPgIqpIiPgDKxIbwk2dnSlAOSvEC2+4a6qBrc7iQxpX3IaKxQJGkXVILP7bIpMcQaZXSsuYBZW5hkGVzXt+9VlcsE6MRMvDSPiPjwgVRsAwk2YBBb52dhcguJrcTVwDFc3TiEjCeH469PbYcN0Dc3H8eTq3mqwUJMi7iFBMHAxnJPomMfz5LQs4/6o31XvFS7oe0+9riP6xA/yyd3RofFpG4illstooR8sZragAfk2bkrhBlmuE0tAPan4y2F/HaCQ/EJTXM55Cbc0MQF1EjzzV4IAKxiABTY2Ibpoer7MGNdt+cO0OV8EkitX8CNciLwbIm1UmE7kmLZGZhCS0XHlRE0pjn+QV+NpGxKPoBpQLBMT2VK3QBsBP479/F5gCa7pdfKvmaxfXBFWrz4TSjyVG3uuiSjUhpITZRITf189Zia3Mb+BxWMmCXWyoRtDMo+UZ5by1dqQ4WhVEXl6nEGMVDnnuAc/QI+6ARB7XQOvZW5COv1Jo1ltwxV35RUKtHKZMjBoCR7o1xTQP+As6k2sQWS++Cmd657cagMfbu1HQx9vZt1IZsXj5MhlCb6TKiauYJGt/Q4EezijBqsMXdUF7uPdUghaxt7mfMqG2RSxjyiMW4pT/gkKdEnZe+VcIYjruFdKYnEtZ/P4JlekEb3nfBGuZ1wM1f0fW1k0OBirLCpLXvtTUotGAF76fgOm6aq922xdYCVhk/ZtTtVE83gAR7hz+GfRIkbWYcif87LUSnWYkEnRn86uk+V6Ke2Gdr52j/KHkClw6SEam7jWJtK7CJlJxyC/+bgo8ILxzSzDg6NVLlDAzgJMhJUoSmqozv44y3jl03gt1xuLcsS1hJ0vHn5V5BJSaxgUbnDoW0qLA1g/7t7lz2clZiYfYyBTCLHpRjzVjAOl29PZpaHig3lMZJTA/iiH3wX3WmjyaRxu9xGjKbSW46SPKXJvaQFXo0j/o7J14/7bD6Xt35sjZhRJx4gdYgmUoLFTcREOEbZqQCGSg2YQNOUip8rntsOkCrq0VWNq2fjPbSUcv0Cbg55CueJBdyBLGFk7m6A0koodpnyu4cBuoO/yIVBFof0r5mUjFp9fJQRRDBL44s1hUASr87dKgEXzZBNgi8cu2VX8Qy16kM4a/qiHotK+3sIXEwkxpdyxbUVUieWkaIWuylJ9syXUuwuWR7j1Hfn4bmo/bHtr7Rqj75uZ9v+rHlEZ0BLmgj/DNeT/wA6CdPGjHogMebbLDVxzaFqTzWEuzvnSU7EU1bYSjGSLOIKUvLsXRRhUbWex+T6vnrUAasdEp5fPBwn7PPe2amBxBj7iJ6A1pbG8xPUa4DvI7v4MaUYXoxi9Lv1puFd4GkKKCE+a/YQA5koxxnpNxLHxkRut2SEIQw3pMmkatpLeMRzX/S3E2NWTSqWa7qVUiTRJKX5TxCHkEJF2MEUVYV+l4z5VAb7nkqRseXqpZIGAgqcqYQMAcNpxJJnF9K22NKKTrV2gTf56ZT2biRQTtfYqAm6pEsV5zQy+BMyBFgHlz9gCMMCHrCSyNiRwnt3vr53H3u37+8AXq+ihvtR1aU4eJ+3TAWOYYRqa2Z1vrCa0MSUHYPqJbLhPgeBv9a+avycC/8IL065VHEnBWGOI88rO2FAScKVsYmKZTG4utrCfXvnwR98VzXzbUgP4C2hkQQzCSPJjbClzO7JCVb/Ao+cGhGb53kXnB3J700L1gJndacMR+VDjpNQKRweHDjJ3VaBQCyJoTY+Inbuql1o8457pf3tHmBpHrAiMVhIssSFYxSp3Tw/9dRwTi/U7uiQNOQmY74ymW+6P9Qk9QPm7kd7yURNLkzTlSb3DldtUFJxFCaKVFR2xfG8++qLnT9Wx1uDzWxlH5KzUM0TNV7OkF+ZHRvAKyIetzFnjn3xk/9cwb3FakG3rxJSxnr55Benqckvhkyv7Gu4JWRJ8Npw1G1spw08HS7iHmfAtv4N15Zfi/fOpsmleMTYl+FKJI2FY1/Zd5ZrYznyUjlLnqfX81ek7WzdMB9UXd8XN1SCyzEknqog0QLV63KP+AIUkExQymlqgAlZFp8Ng9uVP6oUD0jn4CGV5HEbUb6heafpjWZKpuYFIwL9mBBX7ZVM/+CGc3UXRa3WK7S4ftMSluASWqzQZefJWQbfqJWMRVHUZbRpHLj4KPrOV25dL29bKNqPnHkcvYhz5tE52ZLmxgVGK46VzOTtCxSnZl/tLwslp6r9qMUjVoteUolVXpZwUCN/IPdWA6rg3OpRhCA26IxnaOfD7pXi5EdQgNfR17iytL8jo6pISdQDpp60ljox9ST0Br4clpvEKwGZa2PG9VsRhPCxMPF93un7JTtS8gdIzwojaEQYa7SZCLqYMruwMfIgtUwRxK5EeyfryuDX3dONvSyfzGeaxfZSmvCEGC2YjEM6kLudZxkue1XU6GThfGSstLGtyvoxu9y0IvFdbgpL50z/ADlayI3FOI8RW0Dd+f15N1RtU/ijO2DxLmydLDxgdhz6udPVo5OFr3xxrVNIYWImE9jVkxBXskvOKM40p4QJGztWrwEynME8eBG3XQXm6lxz8PVs8lc/YqcpHHGu4zkNodiUyQ+tUASaI49cJtVXPwvCWDg/IB2hAZg63AyTCbZe8hhJ7UyJOJWFy40wJSfZt1JTbnFAOaa29rFZuXDjJBluhB38fG/36sqwhdPEpEmsjiICq1PZWawMYgxk77hdMbb2talQEeVzThp+CK59uo2aQJrxiEqClFYkyi6cI9ue5KYcCGosh2NuYs3iXbW9uKWPkYQ9poPIsGvgwWdFl2sHpVaYSpKEqSCvEfnljw0EPkAuNLa3u24733yuRtqjCt18DGj9cEsYlbYqhQzkhpmS5kbowsC3poIqWZQV/QRBbdfiy+M7QDLjwVCfw62QBu7uiRIZM5lNUMsNR2xiRzVCBJOf26qtwLaxjgo6OOzyufHavVeIZ4Vh/ZgFNsSmF9jg4uXcrE2hmREQ68L4ZoysAQCzC/Wqa8p9mAyqfj77Izp3v78j2UdRSU89pBYjCU3oOmCdT5Uk92oExlHbWhoTWz3y5eya4lDVqI3c4g2scbFUe/TdnBj8CNE5HdPlGY2EckU/oGBlreBccx2XrnX1vqpRZ3vsPZx89Us7tLOSujYPkC+kkC0khMg5ivyu2A+4XBSZEWLg40Xbx11VHzHlq9FL3ZsPNmRCSlRLFwYDoCMqvoaMMCT/GJNd194qbsO7Rmc7PwLdU40SD00LX0N4Xo9Tnqe2c/Wd7VJiQppaWE83LHiCOEoHTAUoPfs6AKWIpERyHi+5HJAFFAS360tZ737dxdCJfvTS5crGIAc3dr5GnM5Lw7IPWVsdakIoH/LdhdRP575qw5qJO0qrkk2UaBavMpWWK5rYpsixQCVyM4KEVDh6Y1JzDKeudcd30sY9ld/+5PH54dTdLqJ8xAJvamIcz/GMGVx6TkX+bcJUEZFkdHhkCbUXLajm3PjJ5m6wZQtPHy5odTzVOOU/I3msr2BpzSmzKhEDBEdiusgdAygzOKSmYsvPd7533YX+2b7vx7kz4K/lRAhJLBw0wuoHEfGhGYYZT349d8vRkRkRWwADEOvVNe87dDt3tfYN/P98N6URpVotqwW2wv1vJG4rjhGylLlJe4yF425jy4aPo5b7hyJryK+rHlKeYaiKQ7tvX+eL4MgDNOwEp5bFE0WGwqOlNNkHjTlhHCXKY4Prx9e6fr1ssB41SKJQwjwk3VFKp9IdgeqsLHc/AmUjzBhcono2R5SLut3MPJNqIY8Q/gvjMYnDBMBU5O9zISIVxjIjbaJBX2O6czlQYVMl/KS5t/lMkUnpb3mJBs5KonXDNAY+nV2JhITlQlGFtoPbdqHWftkMd9PZ0sgoWE/+xLmwlIgdbqkW+eexALpLIuDSxXLnsaqwq3FN1dHV5wFP07zKZ3kp9XpNDbSRSYhIAWwXpJQqOz2dWCZx9skkxh3qc/VW7LAeOp1bn9UWlibJhirNo0U8hkky1SXLroikNRUKQEpsIdWLa3Diseh84333FWeyr/o1k8FHKifrEBbWO22oyOpE14+PKypJfmoVNRdNpFgKE2JZXT357/ODrJ1oRtrFmpGoX5EQ+pFITLAyt2YkY9RYEsYgYwgzaJLX15vNBtfHSNfYPWGPkNgEGJdQtYebiLX07JPs2qL0p4g7cfc5LYozNHPSJ3mEui1gAJayjdkQUvLcVDPFUAN8nKi19vtRDo/TzekBJ04fIf7PiUlcNsbXbQBairqxdo0zTZD+qpjIZvPV1/CQ/anCNW+48hSVyTuUINtXrpozXLh9QDcZp9htnDQlLHb9dG5BcimMYMYmNbZOnT9W5yPK/+3cl/PoxTEN/mwC3tVxJfQBSBxn22nCbCbw03NTizlAzSBMhjNs8nvN0rEAemqx+9NBLlwM1Zfz7l4XAg4b5Q+YhmDaKpmYhlBhKXHujBjJ6mFdSWKpS/sG1/FCXLgiyYblZkXY0jjZz3irz2UfoFxigq6pFAl5AMxySpW7SW/h3HHLY+UEgKL1teIEnrqqeW5vwySBLOcBSQ4gQJvQfxUUY4DMDbiYJoJiXSGWKL9zG9v7a5zVRHZyeQYI1zCxmVEynIrI3n2AG8gNpmCxrUrjLE1bQJgMrI6IcILVpWWrN5pgIiojGmahIkxQE79kuVcrSQmpFzcipqV0Gfz7pH/iWPIMaBmznsyBOSm1LKpjZjbS4JpKyvPL4ClI0EVq0W5d4BZij1zGHlfi3cY7XBawei2OVqhRREi80ceCDnX+ep2SqKA+xjsaW6vbVP1ViuyP0wGRLoS87zD1luIEwawWJM66hoQZL1/us8WNxC0ZypAUv7Hqd21Ic+Yqb1o8ZCxEKJ4gf1oscercozNWQn4DiWm0RHV+qyC1+ZjQaosXcFHTyG9KtXrgFpeH4iaDRA1P4fyayC6qjDoluDw3xjE7+MZjywX1lLG3AAAAMuTDNLSJUq0uASvDBFcsoaaI/TtI+VTu0GapYZYlVbn8L7hCqLlXqLMLVqD/6/8Dzzq44w=="
SEED_MAP=r"""080423=02302642
080287=02301469
080282=02301735
080248=02307519
080502=02302181
080254=02308054
080436=02301929
080117=02308755
080178=02308067
080500=02307620
080251=02310381
080293=02307791
080100=02296823
080053=02298011
080514=02299138
080271=02196272
080036=02301356
080419=02299551
080156=02301428
080045=02301551
080064=02301739
0801094=02199355
080091=02301985
080096=02301526
080256=02285154
080363=02287018
090395=01282004
080198=02290995
079913=21416566
080133=21416566
080476=02292373
080469=02293854
080466=02294262
080467=02294262
080342=02199202
080451=02196629
080353=02199385
080461=02196277
080339=02199160
080403=02187096
080128=01275454
080116=01275788
080106=01275808
080296=01275966
080135=01276046
080378=01276463
080373=01276679
080182=01277089
080387=01277762
080395=01282004
080111=01253007
080071=01253568
080122=01255950
080390=01255949
080431=01256546
0800112=01256571
080487=01257095
080196=01257105
080099=01258983
080314=01259084
080464=01268319
080401=01268364
080170=01268589
080483=01269071
080250=01270759
080336=01271269
079792=01273574
080051=01275014
080174=01275099
080340=01275180
080124=01275233
0800037=01275460
080434=98900104
080040=98900104
080110=99393140
080418=99676706
080162=99753813
080089=99830953
080161=01223079
080101=01223081
080165=01234412
080460=01234548
080067=01234637
080068=01234637
080269=01235865
080074=01236101
080075=01236101
080499=01243192
080107=01243700
079914=82596430
079844=98816734
0799300=98842236
079984=98860521
079919=92342585
079969=97682196
079991=97707761
079954=97732902
079979=97750376
079992=97988645
079967=98223322
079968=98223322
079921=98494675
080004=98611111
080029=98612986
080006=98619043
079933=98619555
079949=98637545
080025=98638831
080002=98647183
079917=98648443
79961=81168264
79942=81247294
799001=81309933
79935=81429740
79907=81561196
79908=81576202
79952=81581780
79958=81617653
79960=81627965
79946=81656654
80019=81690814
79905=81727866
79974=81762435
79999=81890432
79987=81943330
79972=81999746
79944=82060145
79980=82103254
80005=82118536
79985=82164401
79965=82184620
79956=82201663
80013=82304725
79941=82311832
79928=82318692
80028=82405551
80008=82442006
079938=82450071
079939=82450071
79922=82461050
79959=82466226
79988=82477570
79953=82484194
80027=82486703
079924=82551626
079925=82551626
79912=82587621
79836=57768380944
79843=57768182096
79857=57768243626
79861=57768260356
79884=57768282535
79885=57753794576
79889=57768268104
79890=57768381014
79895=57768404291
79835=57751574390
79808=57752118345
79807=57752118345
79806=57752118345
79879=57753207221
79886=57753217636
79860=57753225760
79888=57753229816
79891=57753239012
79864=57753247445
79838=57753287780
79872=57753314052
79894=57762165644
79854=57762225391
79874=57762252175
79855=57762263121
79850=57762311314
79868=57762313694
79848=57762330225
79880=57762575166
79858=57762613515
79847=57762619373
79866=57762626900
79887=57767443224
79837=57767711921
79892=57767868533
79893=57767868533
79842=57768086465
79881=57768281990"""

st.markdown("""<style>
.block-container{padding-top:3.5rem!important;max-width:1450px}
.main-title{font-size:2rem;line-height:1.3;font-weight:800;margin:0 0 .25rem 0;padding:.45rem 0 .1rem}
.sub{color:#667085;margin-bottom:1.2rem}.stMetric{background:#fff;border:1px solid #e5e7eb;border-radius:14px;padding:14px}
[data-testid="stSidebar"]{background:#0b2341}[data-testid="stSidebar"] *{color:white}
</style>""",unsafe_allow_html=True)

def empty_db():
    return {"shipments":{},"payments":{},"nf_map":{},"status":{},"imports":[],"settings":{"client_name":CLIENTE,"cnpj":CNPJ}}

def _github_headers():
    return {"Authorization":f"Bearer {GH_TOKEN}","Accept":"application/vnd.github+json","X-GitHub-Api-Version":"2022-11-28"}

def _remote_get():
    if not (GH_TOKEN and GH_REPO): return None, None
    url=f"https://api.github.com/repos/{GH_REPO}/contents/{GH_DB_PATH}"
    r=requests.get(url,headers=_github_headers(),params={"ref":GH_BRANCH},timeout=20)
    if r.status_code==404: return None, None
    r.raise_for_status(); obj=r.json()
    raw=base64.b64decode(obj["content"]).decode("utf-8")
    return json.loads(raw), obj.get("sha")

def load_db():
    # No Streamlit Cloud, GitHub é a fonte persistente. Local continua funcionando como cache/fallback.
    try:
        remote,_=_remote_get()
        if isinstance(remote,dict):
            with open(DB_FILE,"w",encoding="utf-8") as f: json.dump(remote,f,ensure_ascii=False,indent=2)
            d=remote
        elif os.path.exists(DB_FILE):
            with open(DB_FILE,"r",encoding="utf-8") as f:d=json.load(f)
        else: d=empty_db()
        base=empty_db()
        for k,v in d.items(): base[k]=v
        for k in empty_db():
            if k not in base: base[k]=empty_db()[k]
        return base
    except Exception as e:
        if os.path.exists(DB_FILE):
            try:
                with open(DB_FILE,"r",encoding="utf-8") as f:return json.load(f)
            except: pass
        return empty_db()

def save_db(db):
    """Salva localmente e no GitHub, repetindo a gravação se houver conflito 409 de SHA."""
    raw=json.dumps(db,ensure_ascii=False,indent=2)

    with open(DB_FILE,"w",encoding="utf-8") as f:
        f.write(raw)

    if not (GH_TOKEN and GH_REPO):
        return

    url=f"https://api.github.com/repos/{GH_REPO}/contents/{GH_DB_PATH}"
    last_error=None

    # O GitHub exige o SHA mais recente para atualizar um arquivo existente.
    # Se ocorrer 409, buscamos novamente o SHA atual e repetimos o PUT.
    for attempt in range(4):
        try:
            _,sha=_remote_get()
            payload={
                "message":"Atualiza base FARMATEC",
                "content":base64.b64encode(raw.encode("utf-8")).decode("ascii"),
                "branch":GH_BRANCH
            }
            if sha:
                payload["sha"]=sha

            r=requests.put(url,headers=_github_headers(),json=payload,timeout=25)

            if r.status_code in (200,201):
                return

            if r.status_code==409:
                last_error=RuntimeError("Conflito temporário 409 ao salvar no GitHub.")
                continue

            r.raise_for_status()

        except requests.HTTPError as e:
            last_error=e
            if getattr(e.response,"status_code",None)==409:
                continue
            raise
        except Exception as e:
            last_error=e
            if attempt < 3:
                continue
            raise

    if last_error:
        raise last_error

def digits(v):
    if v is None or (isinstance(v,float) and pd.isna(v)): return ""
    s=str(v).strip()
    if re.fullmatch(r"\d+\.0",s): s=s[:-2]
    return re.sub(r"\D","",s)

def money(v):
    return f"R$ {float(v):,.2f}".replace(",","X").replace(".",",").replace("X",".")

def norm_cnpj(v):
    s=digits(v)
    return s.zfill(14) if len(s)==13 else s

def awb_short(v):
    s=digits(v)
    return s[3:] if len(s)==11 and s.startswith("577") else s

def awb_candidates(v):
    s=digits(v)
    if len(s)==8:return [s,"577"+s]
    if len(s)==11 and s.startswith("577"):return [s,s[3:]]
    return [s]

def to_float(v):
    if pd.isna(v): return 0.0
    if isinstance(v,(int,float)): return float(v)
    s=str(v).replace("R$","").replace(" ","")
    if "," in s:s=s.replace(".","").replace(",",".")
    try:return float(re.sub(r"[^0-9.-]","",s) or 0)
    except:return 0.0

def parse_dt(v):
    if v is None or (isinstance(v,float) and pd.isna(v)): return None
    if isinstance(v,datetime): return v
    if isinstance(v,date): return datetime.combine(v,datetime.min.time())
    if isinstance(v,pd.Timestamp): return v.to_pydatetime()
    # Datas ISO (AAAA-MM-DD) não podem usar dayfirst=True, senão 2026-10-02 vira 10/02/2026.
    txt=str(v).strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}",txt):
        x=pd.to_datetime(txt,errors="coerce",yearfirst=True)
    else:
        x=pd.to_datetime(v,errors="coerce",dayfirst=True)
    return None if pd.isna(x) else x.to_pydatetime()

def seed(db):
    changed=False
    base=json.loads(zlib.decompress(base64.b64decode(BASE_BLOB)).decode())
    for r in base:
        a=r["awb"]
        if a not in db["shipments"]:
            db["shipments"][a]=r; changed=True
    for line in SEED_MAP.splitlines():
        if "=" not in line: continue
        nf,awb=[x.strip() for x in line.split("=",1)]
        if nf and nf not in db["nf_map"]:
            db["nf_map"][nf]={"nf":nf,"awb":digits(awb),"source":"BASE INICIAL"}; changed=True
    if changed: save_db(db)

def nfs_for_awb(db,awb):
    t=awb_short(awb)
    return sorted({v["nf"] for v in db["nf_map"].values() if awb_short(v.get("awb",""))==t})

def find_shipment(db,q):
    t=awb_short(q)
    for s in db["shipments"].values():
        if awb_short(s.get("awb",""))==t:return s
    return None

def payment_dt(p):
    raw=(str(p.get("date",""))+" "+str(p.get("time",""))).strip()
    x=parse_dt(raw)
    return x or datetime.max

def shipment_dt(s):
    return parse_dt(s.get("date","")) or datetime.max

def allocations(db):
    """FIFO: créditos são consumidos por embarques do novo fluxo; uma AWB pode ser rateada."""
    payments=sorted([(k,p) for k,p in db["payments"].items()],key=lambda x:payment_dt(x[1]))
    credits=[{"id":k,"p":p,"remaining":float(p.get("value",0))} for k,p in payments]
    result={}
    new_ship=sorted([s for s in db["shipments"].values() if shipment_dt(s).date()>=CUTOFF],key=shipment_dt)
    ci=0
    for s in new_ship:
        need=float(s.get("total",0)); parts=[]
        while need>0.00001 and ci<len(credits):
            c=credits[ci]
            if c["remaining"]<=0.00001: ci+=1; continue
            use=min(need,c["remaining"])
            parts.append({"payment_id":c["id"],"value":use,"date":c["p"].get("date",""),"time":c["p"].get("time",""),"file":c["p"].get("file","")})
            c["remaining"]-=use; need-=use
            if c["remaining"]<=0.00001: ci+=1
        result[s["awb"]]={"parts":parts,"pending":need}
    return result,credits

def extract_pdf_text(data):
    try:
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    except:return ""

def parse_payment(text):
    clean=re.sub(r"[\t ]+"," ",text or "")
    m=re.search(r"Valor\s*(?:R\$)?\s*([\d.]+,\d{2})",clean,re.I); val=to_float(m.group(1)) if m else 0
    m=re.search(r"Data da transferência\s*(\d{2}/\d{2}/\d{4})",clean,re.I); dt=m.group(1) if m else ""
    m=re.search(r"Efetuada em\s*\d{2}/\d{2}/\d{4}\s*às\s*(\d{2}:\d{2}:\d{2})",clean,re.I); tm=m.group(1) if m else ""
    m=re.search(r"ID da transação\s*[\r\n ]*([A-Z0-9]+)",clean,re.I); tid=m.group(1).upper() if m else ""
    m=re.search(r"Autenticação no comprovante\s*[\r\n ]*([A-Z0-9]+)",clean,re.I); auth=m.group(1).upper() if m else ""
    m=re.search(r"Tipo de Pagamento\s*([^\r\n]+)",text or "",re.I); ptype=m.group(1).strip() if m else ""
    m=re.search(r"Controle\s*[\r\n ]*(\d+)",clean,re.I); control=m.group(1) if m else ""
    return val,dt,tm,tid,auth,ptype,control

def process_excel(data,name,db):
    df=pd.read_excel(io.BytesIO(data))
    required={"AWB","DATA","CODIGOREMETENTE","NOMEREMETENTE","CODIGODESTINATARIO","NOMEDESTINATARIO","CODIGOTOMADOR","TOTAL","ENTREGAPRAZO"}
    miss=required-set(df.columns)
    if miss: raise ValueError("Colunas ausentes: "+", ".join(sorted(miss)))
    f=df[df["CODIGOTOMADOR"].apply(norm_cnpj)==CNPJ].copy()
    new=upd=0; value_new=0
    for _,r in f.iterrows():
        awb=digits(r["AWB"])
        if not awb:continue
        dt=parse_dt(r["DATA"])
        item={"awb":awb,"date":dt.isoformat() if dt else "","sender_cnpj":digits(r["CODIGOREMETENTE"]),
              "sender_name":"" if pd.isna(r["NOMEREMETENTE"]) else str(r["NOMEREMETENTE"]),
              "recipient_doc":digits(r["CODIGODESTINATARIO"]),
              "recipient_name":"" if pd.isna(r["NOMEDESTINATARIO"]) else str(r["NOMEDESTINATARIO"]),
              "total":to_float(r["TOTAL"]),"sla":"" if pd.isna(r["ENTREGAPRAZO"]) else str(r["ENTREGAPRAZO"]),
              "source":name}
        if awb in db["shipments"]:
            upd+=1
            current=db["shipments"][awb]
            # Uma única ficha por AWB: nunca troca informação válida por vazio.
            for k,v in item.items():
                if k in ("total",):
                    if float(v or 0)>0: current[k]=v
                elif v not in (None,""):
                    current[k]=v
            current["source"]=name
        else:
            new+=1;value_new+=item["total"]
            db["shipments"][awb]=item
    db["imports"].append({"file":name,"at":datetime.now().isoformat(timespec="seconds"),"found":len(f),"new":new,"updated":upd,"new_value":value_new})
    save_db(db); return len(f),new,upd,value_new

def add_nf_block(text,db):
    added=existing=0
    for line in text.splitlines():
        awbm=re.search(r"AWB(?:Number)?\s*[:\-]?\s*([\d-]+)",line,re.I)
        if not awbm:
            # aceita tabela AWB TAB NF
            parts=re.split(r"\s+",line.strip())
            if len(parts)>=2 and digits(parts[0]) and digits(parts[1]):
                awb=digits(parts[0]); nfs=re.findall(r"\d+", " ".join(parts[1:]))
            else: continue
        else:
            awb=digits(awbm.group(1))
            left=line[:awbm.start()]
            nfs=re.findall(r"\d+",left)
        for nf in nfs:
            if nf in db["nf_map"] and awb_short(db["nf_map"][nf]["awb"])==awb_short(awb): existing+=1
            else: db["nf_map"][nf]={"nf":nf,"awb":awb,"source":"MANUAL"};added+=1
    save_db(db);return added,existing

def process_status_excel(data,name,db):
    df=pd.read_excel(io.BytesIO(data))
    required={"AWBNumber","StatusDescription","ApproxSLA","Shipper"}
    miss=required-set(df.columns)
    if miss: raise ValueError("Colunas ausentes no relatório de status: "+", ".join(sorted(miss)))
    # O arquivo real traz Shipper como: 06177615000174-NOME DO CLIENTE
    f=df[df["Shipper"].astype(str).apply(lambda x: CNPJ in digits(x))].copy()
    saved=0
    for _,r in f.iterrows():
        num=digits(r["AWBNumber"])
        pref=digits(r.get("AWBPrefix",""))
        # AWBNumber pode perder zeros à esquerda no Excel. Primeiro tenta localizar na base financeira.
        matches=[x for x in db["shipments"].values() if awb_short(x.get("awb",""))==num or awb_short(x.get("awb","")).lstrip("0")==num.lstrip("0")]
        if matches: awb=matches[0]["awb"]
        else:
            short=num.zfill(8) if len(num)<=8 else num
            awb=(pref+short) if pref and not short.startswith(pref) else short
        sla=r.get("ApproxSLA","")
        sla_dt=parse_dt(sla)
        exec_dt=parse_dt(r.get("ExecutionDateTime",""))
        shipper="" if pd.isna(r.get("Shipper","")) else str(r.get("Shipper","")).strip()
        consignee="" if pd.isna(r.get("Consignee","")) else str(r.get("Consignee","")).strip()
        # Remove o CNPJ do começo do nome do remetente quando o relatório vier como CNPJ-NOME.
        sender_status=re.sub(r"^\s*0?6177615000174\s*[-–—:]?\s*","",shipper,flags=re.I).strip() or shipper
        db["status"][awb_short(awb)]={"awb":awb,"status":"" if pd.isna(r["StatusDescription"]) else str(r["StatusDescription"]),"sla_date":sla_dt.isoformat() if sla_dt else "","updated_event":exec_dt.isoformat() if exec_dt else "","shipper":shipper,"sender_name":sender_status,"recipient_name":consignee,"source":name,"imported_at":datetime.now().isoformat(timespec="seconds")}
        # Atualiza a ficha central da AWB. Todos os menus e relatórios leem esta mesma ficha.
        central=find_shipment(db,awb)
        if central is None:
            central={"awb":awb,"date":"","sender_cnpj":"","sender_name":"","recipient_doc":"","recipient_name":"","total":0.0,"sla":"","source":"STATUS"}
            db["shipments"][awb]=central
        if not str(central.get("sender_name","")).strip() and sender_status:
            central["sender_name"]=sender_status
        if not str(central.get("recipient_name","")).strip() and consignee:
            central["recipient_name"]=consignee
        central["status"]="" if pd.isna(r["StatusDescription"]) else str(r["StatusDescription"])
        central["delivery_forecast"]=sla_dt.isoformat() if sla_dt else ""
        central["status_source"]=name
        saved+=1
    db["imports"].append({"file":name,"at":datetime.now().isoformat(timespec="seconds"),"type":"STATUS","found":len(f),"new":saved,"updated":0,"new_value":0})
    save_db(db);return len(f),saved

def status_for_awb(db,awb):
    t=awb_short(awb)
    if t in db.get("status",{}): return db["status"][t]
    for k,v in db.get("status",{}).items():
        if k.lstrip("0")==t.lstrip("0"): return v
    return None

def unified_shipment(db,s):
    """Retorna a ficha única da AWB, complementando campos vazios com o relatório de status."""
    u=dict(s)
    stt=status_for_awb(db,u.get("awb","")) or {}
    if not str(u.get("sender_name","")).strip(): u["sender_name"]=stt.get("sender_name","")
    if not str(u.get("recipient_name","")).strip(): u["recipient_name"]=stt.get("recipient_name","")
    u["status"]=stt.get("status",u.get("status","SEM STATUS"))
    u["delivery_forecast"]=stt.get("sla_date",u.get("delivery_forecast",""))
    return u

def norm_text(v):
    return "".join(c for c in unicodedata.normalize("NFD",str(v or "").lower()) if unicodedata.category(c)!="Mn")

def is_delivered(status):
    return norm_text(status).strip() in {"entregue","delivered"}

def is_overdue(db,awb,now=None):
    stt=status_for_awb(db,awb) or {}
    due=parse_dt(stt.get("sla_date",""))
    if not due or is_delivered(stt.get("status","")): return False
    now=now or datetime.now()
    return due < now

def status_excel(db, only_pending=False, only_overdue=False):
    wb=Workbook();ws=wb.active;ws.title="Status FARMATEC"
    headers=["AWB","NF","Status","Previsão/SLA","Data emissão","Remetente","Destinatário","Valor frete","Pagamento"]
    ws.append(headers)
    alloc,_=allocations(db)
    for s in sorted(db["shipments"].values(),key=shipment_dt):
        s=unified_shipment(db,s)
        stt=status_for_awb(db,s["awb"])
        status=s.get("status","SEM STATUS")
        if only_pending and is_delivered(status): continue
        if only_overdue and not is_overdue(db,s["awb"]): continue
        dt=shipment_dt(s); info=alloc.get(s["awb"])
        pay="CONTA CORRENTE AZUL" if dt.date()<CUTOFF else (" / ".join(f'{x["date"]} {x["time"]}'.strip() for x in info["parts"]) if info and info["parts"] else "PENDENTE")
        sla=parse_dt((stt or {}).get("sla_date",""))
        ws.append([s["awb"],", ".join(nfs_for_awb(db,s["awb"])) or "-",status,sla.strftime("%d/%m/%Y %H:%M") if sla else "",dt.strftime("%d/%m/%Y %H:%M"),s.get("sender_name",""),s.get("recipient_name",""),s.get("total",0),pay])
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="0B2341")
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(max(len(str(x.value or "")) for x in col)+2,38)
    out=io.BytesIO();wb.save(out);return out.getvalue()

def rows_for_report(db,payment_id=None):
    alloc,_=allocations(db); rows=[]
    for s in sorted(db["shipments"].values(),key=shipment_dt):
        s=unified_shipment(db,s)
        dt=shipment_dt(s)
        old=dt.date()<CUTOFF
        nfs=", ".join(nfs_for_awb(db,s["awb"])) or "-"
        if old:
            if payment_id is not None: continue
            rows.append({**s,"nf":nfs,"payment":"CONTA CORRENTE AZUL","payment_date":"","payment_time":"","allocated":0,"pending":0})
        else:
            info=alloc.get(s["awb"],{"parts":[],"pending":float(s.get("total",0))})
            if payment_id is None:
                desc=" / ".join([f'{p["date"]} {p["time"]}'.strip() for p in info["parts"]]) or "PENDENTE"
                rows.append({**s,"nf":nfs,"payment":desc,"payment_date":"","payment_time":"","allocated":sum(p["value"] for p in info["parts"]),"pending":info["pending"]})
            else:
                for p in info["parts"]:
                    if p["payment_id"]==payment_id:
                        rows.append({**s,"nf":nfs,"payment":p["file"],"payment_date":p["date"],"payment_time":p["time"],"allocated":p["value"],"pending":info["pending"]})
    return rows

def make_pdf(db,payment_id):
    p=db["payments"][payment_id]; rows=rows_for_report(db,payment_id)
    used=sum(r["allocated"] for r in rows); received=float(p["value"]); remain=received-used
    bio=io.BytesIO(); doc=SimpleDocTemplate(bio,pagesize=landscape(A4),rightMargin=22,leftMargin=22,topMargin=22,bottomMargin=22)
    ss=getSampleStyleSheet(); story=[]
    title=ParagraphStyle("t",parent=ss["Title"],alignment=TA_CENTER,fontSize=22,leading=27)
    story += [Spacer(1,40),Paragraph("FARMATEC",title),Paragraph("Relatório de utilização de saldo",ss["Heading2"]),Spacer(1,20)]
    cover=[["Comprovante",p.get("file","")],["Data / hora",f'{p.get("date","")} às {p.get("time","")}'.replace(" às "," às ")],
           ["Tipo de pagamento",p.get("payment_type","")],["ID da transação",p.get("transaction_id","")],
           ["Autenticação",p.get("authentication","")],["Controle",p.get("control","")],
           ["Valor recebido",money(received)],["Fretes vinculados",money(used)],["Saldo do comprovante",money(remain)],["Quantidade de AWBs",str(len({r["awb"] for r in rows}))]]
    t=Table(cover,colWidths=[150,360]);t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.5,colors.grey),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#EAF2F8")),("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),("PADDING",(0,0),(-1,-1),8)]));story+=[t,Spacer(1,20)]
    if remain<0: story.append(Paragraph(f"Valor pendente: {money(abs(remain))}",ss["Heading2"]))
    else: story.append(Paragraph(f"Saldo remanescente: {money(remain)}",ss["Heading2"]))
    story += [PageBreak(),Paragraph("Embarques vinculados ao comprovante",ss["Heading2"]),Spacer(1,8)]
    data=[["AWB","NF","Data","CNPJ remet.","Remetente","CPF/CNPJ dest.","Destinatário","Frete","SLA","Valor deste comprov."]]
    for r in rows:
        dt=shipment_dt(r)
        data.append([r["awb"],r["nf"],dt.strftime("%d/%m/%Y %H:%M") if dt!=datetime.max else "",r.get("sender_cnpj",""),str(r.get("sender_name",r.get("shipper_name","")))[:24],r.get("recipient_doc",""),str(r.get("recipient_name",""))[:24],money(r["total"]),str(r.get("sla","") or (parse_dt(r.get("delivery_forecast","" )).strftime("%d/%m/%Y %H:%M") if parse_dt(r.get("delivery_forecast","")) else "")),money(r["allocated"])])
    tab=Table(data,repeatRows=1,colWidths=[70,65,82,82,110,82,110,65,38,78])
    tab.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#0B2341")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),6.5),("GRID",(0,0),(-1,-1),.25,colors.lightgrey),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),3)]))
    story.append(tab);doc.build(story);return bio.getvalue()

def make_excel(db,payment_id):
    p=db["payments"][payment_id];rows=rows_for_report(db,payment_id)
    wb=Workbook();ws=wb.active;ws.title="Resumo"
    used=sum(r["allocated"] for r in rows)
    summary=[["FARMATEC - RELATÓRIO DE UTILIZAÇÃO DE SALDO"],["Comprovante",p.get("file","")],["Data",p.get("date","")],["Hora",p.get("time","")],["Tipo de pagamento",p.get("payment_type","")],["ID da transação",p.get("transaction_id","")],["Autenticação",p.get("authentication","")],["Controle",p.get("control","")],["Valor recebido",float(p["value"])],["Fretes vinculados",used],["Saldo",float(p["value"])-used]]
    for row in summary:ws.append(row)
    ws2=wb.create_sheet("Embarques")
    headers=["AWB","NF","Data","CNPJ remetente","Nome remetente","CPF/CNPJ destinatário","Nome destinatário","Valor frete","SLA dias úteis","Comprovante","Data comprovante","Hora comprovante","Valor utilizado"]
    ws2.append(headers)
    for r in rows:
        ws2.append([r["awb"],r["nf"],shipment_dt(r).strftime("%d/%m/%Y %H:%M"),r.get("sender_cnpj",""),r.get("sender_name",r.get("shipper_name","")),r.get("recipient_doc",""),r.get("recipient_name",""),r["total"],r.get("sla","") or (parse_dt(r.get("delivery_forecast","" )).strftime("%d/%m/%Y %H:%M") if parse_dt(r.get("delivery_forecast","")) else ""),r["payment"],r["payment_date"],r["payment_time"],r["allocated"]])
    for c in ws2[1]:
        c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="0B2341");c.alignment=Alignment(horizontal="center")
    for col in ws2.columns:
        letter=col[0].column_letter;ws2.column_dimensions[letter].width=min(max(len(str(x.value or "")) for x in col)+2,35)
    out=io.BytesIO();wb.save(out);return out.getvalue()

db=load_db();seed(db)
# Compatibilidade com registros criados por versões anteriores
_db_changed=False
for _awb,_s in db.get("shipments",{}).items():
    defaults={"awb":_awb,"date":"","sender_cnpj":"","sender_name":"","recipient_doc":"","recipient_name":"","total":0.0,"sla":"","source":"BASE ANTERIOR"}
    for _k,_v in defaults.items():
        if _k not in _s:
            _s[_k]=_v; _db_changed=True
# Migração: informações já importadas em Status passam a completar a ficha central.
for _awb,_s in db.get("shipments",{}).items():
    _st=status_for_awb(db,_awb) or {}
    if not str(_s.get("sender_name","")).strip() and _st.get("sender_name"):
        _s["sender_name"]=_st["sender_name"]; _db_changed=True
    if not str(_s.get("recipient_name","")).strip() and _st.get("recipient_name"):
        _s["recipient_name"]=_st["recipient_name"]; _db_changed=True
    if _st.get("status"):
        _s["status"]=_st["status"]; _db_changed=True
    if _st.get("sla_date"):
        _s["delivery_forecast"]=_st["sla_date"]; _db_changed=True
if _db_changed: save_db(db)
alloc,credits=allocations(db)
all_ship=list(db["shipments"].values())
new_ship=[s for s in all_ship if shipment_dt(s).date()>=CUTOFF]
old_ship=[s for s in all_ship if shipment_dt(s).date()<CUTOFF]
total_paid=sum(float(p.get("value",0)) for p in db["payments"].values())
new_freight=sum(float(s.get("total",0)) for s in new_ship)
balance=total_paid-new_freight

with st.sidebar:
    st.markdown("## 💊 FARMATEC")
    page=st.radio("Menu",["Dashboard","Importar embarques","Comprovantes","NF × AWB","Status/Acompanhamento","Consulta","Relatórios","Histórico/Backup"])
    st.caption("Novo fluxo de saldo desde 22/09/2026")
    st.code("CNPJ 06.177.615/0001-74")
    st.caption("💾 Persistência GitHub ativa" if (GH_TOKEN and GH_REPO) else "⚠️ Persistência somente local")
    if st.button("Sair", use_container_width=True):
        st.session_state["farmatec_logged_in"]=False; st.rerun()

st.markdown('<div class="main-title">FARMATEC — Controle de saldo</div>',unsafe_allow_html=True)
st.markdown('<div class="sub">Conciliação de pagamentos, fretes, NF × AWB e relatórios do cliente.</div>',unsafe_allow_html=True)

if page=="Dashboard":
    c1,c2,c3,c4=st.columns(4)
    c1.metric("💰 Total pago",money(total_paid));c2.metric("✈️ Fretes desde 22/09",money(new_freight));c3.metric("✅ Saldo",money(balance));c4.metric("📦 AWBs novo fluxo",len(new_ship))
    st.caption(f"Base histórica: {len(all_ship)} AWBs | anteriores a 22/09: {len(old_ship)} (CONTA CORRENTE AZUL)")
    if balance<0:st.error(f"🔴 Valor pendente para continuidade: {money(abs(balance))}")
    elif total_paid:st.success(f"🟢 Saldo disponível: {money(balance)}")
    show=[]
    for s in sorted(all_ship,key=shipment_dt,reverse=True)[:100]:
        dt=shipment_dt(s); info=alloc.get(s["awb"])
        origem="CONTA CORRENTE AZUL" if dt.date()<CUTOFF else (" / ".join(f'{x["date"]} {x["time"]}'.strip() for x in info["parts"]) if info and info["parts"] else "PENDENTE")
        show.append({"AWB":s["awb"],"NF":", ".join(nfs_for_awb(db,s["awb"])) or "-","Data":dt.strftime("%d/%m/%Y %H:%M"),"Frete":s["total"],"Pagamento":origem})
    st.dataframe(pd.DataFrame(show),use_container_width=True,hide_index=True,column_config={"Frete":st.column_config.NumberColumn(format="R$ %.2f")})

elif page=="Importar embarques":
    st.subheader("📥 Importar/atualizar base de embarques")
    st.info("Pode subir relatórios com períodos antigos e novos. Antes de 22/09/2026 o embarque fica como CONTA CORRENTE AZUL e não consome os comprovantes.")
    files=st.file_uploader("Planilhas Excel",type=["xlsx","xls"],accept_multiple_files=True)
    if st.session_state.get("import_done"):
        st.success(st.session_state.pop("import_done"))
    if st.button("Processar planilhas",type="primary",disabled=not files):
        msgs=[]
        for f in files:
            try:
                found,new,upd,val=process_excel(f.getvalue(),f.name,db)
                msgs.append(f"{f.name}: {found} FARMATEC | {new} novas AWBs | {upd} atualizadas | {money(val)} acrescentados")
            except Exception as e: st.error(str(e))
        if msgs:
            st.session_state["import_done"]="✅ Importação concluída com sucesso! Base de embarques atualizada.\n\n"+" | ".join(msgs)
            st.rerun()

elif page=="Comprovantes":
    st.subheader("💳 Comprovantes")
    pdf=st.file_uploader("Comprovante PDF",type=["pdf"])
    if pdf:
        data=pdf.getvalue();text=extract_pdf_text(data);val,dt,tm,tid,auth,ptype,control=parse_payment(text)
        c1,c2,c3=st.columns(3)
        val=c1.number_input("Valor",min_value=0.0,value=float(val),format="%.2f")
        dt=c2.text_input("Data",value=dt);tm=c3.text_input("Hora",value=tm)
        if st.button("Salvar comprovante",type="primary"):
            h=hashlib.sha256(data).hexdigest()
            dup_key=None
            if h in db["payments"]: dup_key=h
            if dup_key is None:
                for _k,_x in db["payments"].items():
                    if (tid and _x.get("transaction_id")==tid) or (auth and _x.get("authentication")==auth): dup_key=_k; break
            if dup_key:
                # Não duplica saldo, mas completa cadastros antigos com os dados lidos do PDF.
                oldp=db["payments"][dup_key]
                for _k,_v in {"file":pdf.name,"date":dt,"time":tm,"transaction_id":tid,"authentication":auth,"payment_type":ptype,"control":control}.items():
                    if _v: oldp[_k]=_v
                if val>0: oldp["value"]=val
                save_db(db);st.session_state["payment_msg"]="⚠️ Comprovante já cadastrado. O saldo não foi duplicado e os dados da transação foram atualizados.";st.rerun()
            else:
                db["payments"][h]={"file":pdf.name,"value":val,"date":dt,"time":tm,"transaction_id":tid,"authentication":auth,"payment_type":ptype,"control":control,"created_at":datetime.now().isoformat(timespec="seconds")}
                save_db(db);st.session_state["payment_msg"]="✅ Comprovante salvo com sucesso.";st.rerun()
    if st.session_state.get("payment_msg"): st.success(st.session_state.pop("payment_msg"))
    if db["payments"]:
        data=[]
        _,cr=allocations(db)
        rem={x["id"]:x["remaining"] for x in cr}
        for k,p in sorted(db["payments"].items(),key=lambda x:payment_dt(x[1])):
            data.append({"Data":p.get("date",""),"Hora":p.get("time",""),"ID transação":p.get("transaction_id",""),"Arquivo":p.get("file",""),"Valor":p["value"],"Saldo do crédito":rem.get(k,p["value"])})
        st.dataframe(pd.DataFrame(data),use_container_width=True,hide_index=True,column_config={"Valor":st.column_config.NumberColumn(format="R$ %.2f"),"Saldo do crédito":st.column_config.NumberColumn(format="R$ %.2f")})

elif page=="NF × AWB":
    st.subheader("🧾 Base NF × AWB")
    pending=[]
    for sh in sorted(db["shipments"].values(),key=shipment_dt,reverse=True):
        if not nfs_for_awb(db,sh.get("awb","")):
            u=unified_shipment(db,sh); dt=shipment_dt(u)
            pending.append({"AWB":u.get("awb",""),"NF":"","Data emissão":dt.strftime("%d/%m/%Y %H:%M") if dt!=datetime.max else "","Destinatário":u.get("recipient_name",""),"Valor frete":float(u.get("total",0))})
    st.caption(f"{len(db['nf_map'])} NFs cadastradas | {len(pending)} AWBs ainda sem NF vinculada")
    txt=st.text_area("Adicionar novos vínculos em massa",height=160,placeholder="NF 080423 — AWB 02302642")
    if st.button("Salvar vínculos em massa",type="primary",disabled=not txt.strip()):
        a,e=add_nf_block(txt,db);st.session_state["nf_msg"]=f"✅ {a} novos vínculos salvos | {e} já existentes.";st.rerun()
    if st.session_state.get("nf_msg"): st.success(st.session_state.pop("nf_msg"))
    if pending:
        st.markdown("### ⚠️ AWBs sem NF vinculada")
        edited=st.data_editor(pd.DataFrame(pending),hide_index=True,use_container_width=True,disabled=["AWB","Data emissão","Destinatário","Valor frete"],column_config={"NF":st.column_config.TextColumn("NF",help="Digite a NF e salve. Para várias NFs, separe por / ou vírgula."),"Valor frete":st.column_config.NumberColumn(format="R$ %.2f")},key="pending_nf_editor")
        if st.button("Salvar NFs preenchidas",type="primary"):
            count=0
            for _,r in edited.iterrows():
                val=str(r.get("NF","")).strip()
                if not val: continue
                for nf in re.findall(r"\d+",val):
                    db["nf_map"][nf]={"nf":nf,"awb":digits(r["AWB"]),"source":"PENDÊNCIA MANUAL"};count+=1
            if count:
                save_db(db);st.session_state["nf_msg"]=f"✅ {count} vínculo(s) de NF salvo(s).";st.rerun()
            else: st.warning("Preencha pelo menos uma NF.")
    st.markdown("### Base vinculada")
    rows=sorted(db["nf_map"].values(),key=lambda x:x["nf"],reverse=True)
    st.dataframe(pd.DataFrame(rows)[["nf","awb"]].rename(columns={"nf":"NF","awb":"AWB"}),use_container_width=True,hide_index=True)

elif page=="Status/Acompanhamento":
    st.subheader("🚚 Status e acompanhamento")
    st.info("Importe o relatório rptAWBStatus. O app filtra somente o Shipper com CNPJ 06.177.615/0001-74 e concilia pela AWB. Esta importação não altera o saldo financeiro.")
    sf=st.file_uploader("Relatório de status Excel",type=["xlsx","xls"],key="status_upload")
    if st.button("Importar status",type="primary",disabled=not sf):
        try:
            found,saved=process_status_excel(sf.getvalue(),sf.name,db)
            st.session_state["status_done"]=f"✅ Importação de status concluída! {found} registros FARMATEC encontrados | {saved} status processados.";st.rerun()
        except Exception as e: st.error(str(e))
    if st.session_state.get("status_done"): st.success(st.session_state.pop("status_done"))
    rows=[]
    for sh in sorted(db["shipments"].values(),key=shipment_dt,reverse=True):
        sh=unified_shipment(db,sh)
        x=status_for_awb(db,sh["awb"]); sla=parse_dt((x or {}).get("sla_date",""))
        rows.append({"AWB":sh["awb"],"NF":", ".join(nfs_for_awb(db,sh["awb"])) or "-","Status":(x or {}).get("status","SEM STATUS"),"Previsão/SLA":sla.strftime("%d/%m/%Y %H:%M") if sla else "","Data emissão":shipment_dt(sh).strftime("%d/%m/%Y %H:%M"),"Destinatário":sh.get("recipient_name","")})
    rdf=pd.DataFrame(rows)
    if not rdf.empty:
        statuses=["Todos","🔴 FORA DO PRAZO"]+sorted(rdf["Status"].dropna().unique().tolist())
        filt=st.selectbox("Filtrar status",statuses)
        if filt=="Todos": view=rdf
        elif filt=="🔴 FORA DO PRAZO": view=rdf[rdf["AWB"].apply(lambda a:is_overdue(db,a))]
        else: view=rdf[rdf["Status"]==filt]
        st.dataframe(view,use_container_width=True,hide_index=True)
        a,b,c=st.columns(3)
        a.download_button("⬇️ Relatório completo",status_excel(db,False),file_name="FARMATEC_status_completo.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
        b.download_button("⬇️ Somente não entregues",status_excel(db,True),file_name="FARMATEC_status_pendentes.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
        c.download_button("⬇️ Fora do prazo",status_excel(db,False,True),file_name="FARMATEC_fora_do_prazo.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)

elif page=="Consulta":
    st.subheader("🔎 Consulta por NF ou AWB")
    q=st.text_input("Digite NF, AWB curta ou AWB completa")
    if q:
        qd=digits(q); maps=[]
        if qd in db["nf_map"]:maps=[db["nf_map"][qd]]
        else:maps=[v for v in db["nf_map"].values() if awb_short(v["awb"])==awb_short(qd)]
        awb=maps[0]["awb"] if maps else qd
        s=find_shipment(db,awb)
        if s:
            dt=shipment_dt(s); info=alloc.get(s["awb"])
            pagamento="CONTA CORRENTE AZUL" if dt.date()<CUTOFF else (" / ".join(f'{x["date"]} {x["time"]} ({money(x["value"])})' for x in info["parts"]) if info and info["parts"] else "PENDENTE")
            c1,c2,c3,c4=st.columns(4);c1.metric("AWB",s["awb"]);c2.metric("NF(s)",", ".join(nfs_for_awb(db,s["awb"])) or "-");c3.metric("Frete",money(s["total"]));c4.metric("SLA",f'{s["sla"]} dias úteis')
            stt=status_for_awb(db,s["awb"]); sla_status=parse_dt((stt or {}).get("sla_date",""))
            st.write(f"**Data:** {dt.strftime('%d/%m/%Y %H:%M')}  \n**Remetente:** {s.get('sender_name',s.get('shipper_name',''))} — {s.get('sender_cnpj','')}  \n**Destinatário:** {s.get('recipient_name','')} — {s.get('recipient_doc','')}  \n**Pagamento:** {pagamento}  \n**Status atual:** {(stt or {}).get('status','SEM STATUS')}  \n**Previsão/SLA de entrega:** {sla_status.strftime('%d/%m/%Y %H:%M') if sla_status else '-'}")
        elif maps:st.warning(f"Vínculo encontrado: NF(s) {', '.join(x['nf'] for x in maps)} → AWB {awb}, mas ainda sem dados do relatório de frete.")
        else:st.error("Nenhum registro encontrado.")

elif page=="Relatórios":
    st.subheader("📄 Relatório por comprovante")
    if not db["payments"]:st.info("Cadastre um comprovante primeiro.")
    else:
        opts=sorted(db["payments"].items(),key=lambda x:payment_dt(x[1]))
        labels={k:f'{p.get("date","")} {p.get("time","")} — {money(p["value"])} — {p.get("file","")}' for k,p in opts}
        pid=st.selectbox("Selecione o comprovante",[k for k,_ in opts],format_func=lambda k:labels[k])
        rows=rows_for_report(db,pid);p=db["payments"][pid];used=sum(r["allocated"] for r in rows);saldo=float(p["value"])-used
        c1,c2,c3,c4=st.columns(4);c1.metric("Valor recebido",money(p["value"]));c2.metric("Fretes vinculados",money(used));c3.metric("Saldo",money(saldo));c4.metric("AWBs",len({r["awb"] for r in rows}))
        if saldo<0:st.error(f"Valor pendente: {money(abs(saldo))}")
        preview=pd.DataFrame([{"AWB":r["awb"],"NF":r["nf"],"Data":shipment_dt(r).strftime("%d/%m/%Y %H:%M"),"Remetente":r.get("sender_name",r.get("shipper_name","")),"Destinatário":r.get("recipient_name",""),"Frete":r["total"],"SLA":r.get("sla","") or (parse_dt(r.get("delivery_forecast","" )).strftime("%d/%m/%Y %H:%M") if parse_dt(r.get("delivery_forecast","")) else ""),"Valor deste comprovante":r["allocated"]} for r in rows])
        st.dataframe(preview,use_container_width=True,hide_index=True,column_config={"Frete":st.column_config.NumberColumn(format="R$ %.2f"),"Valor deste comprovante":st.column_config.NumberColumn(format="R$ %.2f")})
        a,b=st.columns(2)
        a.download_button("⬇️ Baixar relatório PDF + capa",make_pdf(db,pid),file_name=f"FARMATEC_relatorio_{p.get('date','').replace('/','-')}.pdf",mime="application/pdf",use_container_width=True)
        b.download_button("⬇️ Baixar relatório Excel",make_excel(db,pid),file_name=f"FARMATEC_relatorio_{p.get('date','').replace('/','-')}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)

elif page=="Histórico/Backup":
    st.subheader("🗂️ Histórico de importações")
    if db["imports"]:st.dataframe(pd.DataFrame(db["imports"]),use_container_width=True,hide_index=True)
    st.subheader("💾 Backup")
    raw=json.dumps(db,ensure_ascii=False,indent=2).encode()
    st.download_button("⬇️ Baixar backup JSON",raw,file_name=f"farmatec_backup_{datetime.now():%Y%m%d_%H%M}.json",mime="application/json")
    up=st.file_uploader("Restaurar backup",type=["json"])
    if up and st.button("Restaurar"):
        try:
            nd=json.loads(up.getvalue().decode())
            with open(DB_FILE,"w",encoding="utf-8") as f:json.dump(nd,f,ensure_ascii=False,indent=2)
            st.success("Backup restaurado.");st.rerun()
        except Exception as e:st.error(str(e))
